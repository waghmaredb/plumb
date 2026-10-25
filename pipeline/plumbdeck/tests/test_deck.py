"""Rendering, lint and the real GEH run."""
import json
import re
import shutil

import pytest
from pptx import Presentation

from plumbdeck.build import build_deck, write_outputs
from plumbdeck.cli import main
from plumbdeck.interlock import BANNER
from plumbdeck.lint import find_leaks, lint_dir, parse_markdown, parse_pptx, title_words, withheld_fingerprints
from plumbdeck.spec import SpecError, load_spec

from conftest import RUN3, RUN4, SCORECARD, SPEC, make_run, make_signoff

GEH_REASONS = "internal profile; 1 attachment unreconciled; resolution 0.2; no QA gate 3 sign-off"


@pytest.fixture(scope="module")
def geh(tmp_path_factory):
    out = tmp_path_factory.mktemp("geh3")
    deck = build_deck(SPEC, RUN3)
    rep = write_outputs(deck, out)
    return deck, out, rep


# ------------------------------------------------------------ real GEH run 3
def test_geh_run3_is_draft_with_expected_reasons(geh):
    deck, out, rep = geh
    assert deck.draft
    assert "; ".join(deck.interlock.reasons) == GEH_REASONS
    assert sorted(p.name for p in out.iterdir()) == ["build.json", "deck-DRAFT.md", "deck-DRAFT.pptx", "lint.json", "titles.md"]
    assert rep["ok"], rep["errors"]
    assert {w["check"] for w in rep["warnings"]} == {"placeholders"}


def test_geh_pptx_banner_on_every_slide(geh):
    deck, out, _ = geh
    prs = Presentation(str(out / "deck-DRAFT.pptx"))
    assert abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.01
    assert len(prs.slides) >= len(deck.slides)
    for sl in prs.slides:
        banners = [sh for sh in sl.shapes if sh.name == "plumbdeck-banner"]
        assert len(banners) == 1
        assert banners[0].text_frame.text == f"{BANNER} ({GEH_REASONS})"
        footer = [sh for sh in sl.shapes if sh.name == "plumbdeck-footer"][0].text_frame.text
        assert footer.startswith("DRAFT — not for board use. Deliverable")


def test_geh_markdown_matches_pptx(geh):
    deck, out, _ = geh
    mds = parse_markdown((out / "deck-DRAFT.md").read_text(encoding="utf-8"))
    pps = parse_pptx(out / "deck-DRAFT.pptx")
    assert [s["id"] for s in mds] == [s.id for s in deck.slides]
    assert [s["id"] for s in mds] == list(dict.fromkeys(p["id"] for p in pps))
    first = {p["id"]: p for p in pps if p["page"] == 1}
    for s in mds:
        assert first[s["id"]]["title"] == s["title"]
        assert s["banner"] and GEH_REASONS in s["banner"]
        if s["sources"]:
            assert first[s["id"]]["source"] == s["sources"][0]


def test_geh_s02_draft_disclosure(geh):
    deck, out, _ = geh
    s02 = next(s for s in parse_markdown((out / "deck-DRAFT.md").read_text()) if s["id"] == "S02")
    assert s02["title"] == "DRAFT — not for board use: 10 claims released, 4 withheld and listed"
    assert "Verification performed" not in s02["text"]
    assert "Released: 10 claims. Withheld: 4 claims (listed)." in s02["text"]
    assert "0 of 1" in s02["text"] and "| internal |" in s02["text"]


def test_geh_withheld_peer_slide_becomes_placeholder(geh):
    deck, out, _ = geh
    s06 = next(s for s in deck.slides if s.id == "S06")
    assert s06.kind == "placeholder" and s06.placeholder_reason == "withheld"
    md = (out / "deck-DRAFT.md").read_text()
    assert "Withheld — see gap register G-1, G-2, G-3, G-4" in md
    s09 = next(s for s in deck.slides if s.id == "S09")
    assert s09.kind == "data" and "level 4 of 5" in s09.title
    assert s09.source_line == ("Source: T2 Great Eastern Holdings Limited annual report FY2025 pp. 8, 14, 18, 31, 40, 67, 229 "
                               "(attached; publisher hash unreconciled), 2026-10-25; confidence: model-inferred")


def test_geh_no_withheld_figures_anywhere(geh):
    deck, out, _ = geh
    text = (out / "deck-DRAFT.md").read_text() + "\n".join(p["text"] for p in parse_pptx(out / "deck-DRAFT.pptx"))
    for fig in ("66.3", "44.7", "43.4", "345,423", "S$16"):
        assert fig not in text


def test_titles_export_is_horizontal_flow(geh):
    deck, out, _ = geh
    t = (out / "titles.md").read_text()
    rows = [l for l in t.splitlines() if re.match(r"^\| (S|A)\d", l)]
    assert len(rows) == len(deck.slides) == 41
    assert GEH_REASONS in t


# ---------------------------------------------------------------- final deck
def test_final_deck_has_statement_and_no_banner(tmp):
    run = make_run(tmp)
    deck = build_deck(SPEC, run, scorecard=SCORECARD, signoff=make_signoff(tmp, run))
    rep = write_outputs(deck, tmp / "out")
    assert not deck.draft
    prs = Presentation(str(tmp / "out" / "deck.pptx"))
    for sl in prs.slides:
        assert not [sh for sh in sl.shapes if sh.name == "plumbdeck-banner"]
        footer = [sh for sh in sl.shapes if sh.name == "plumbdeck-footer"][0].text_frame.text
        assert "All citations fetched and verified" in footer
    md = (tmp / "out" / "deck.md").read_text()
    assert "Verification performed." in md and BANNER not in md
    assert "Reviewed and approved by: Engagement Partner (fictional), Lead partner on 2026-10-26" in md
    # a final deck with placeholders still fails lint: not board-ready, but only for that reason
    assert {e["check"] for e in rep["errors"]} == {"placeholders"}
    s02 = next(s for s in deck.slides if s.id == "S02")
    assert s02.title == "This deck is AI-assisted, human-accountable, and every citation was verified before release"


def test_scorecard_slide_two_axis(tmp):
    deck = build_deck(SPEC, RUN3, scorecard=SCORECARD)
    sc = json.loads(SCORECARD.read_text())
    s08 = next(s for s in deck.slides if s.id == "S08")
    assert s08.kind == "data"
    assert s08.title.startswith(f"{sc['org_label']} scores {sc['capability_score_rounded']:g} of 100, {sc['band']}")
    assert title_words(s08.title) <= 15
    md = write_outputs(deck, tmp / "o") and (tmp / "o" / "deck-DRAFT.md").read_text()
    assert "Self-view (uncapped)" in md and "Evidenced view (cap rule)" in md
    assert s08.source_line.startswith("Source: T1 ") and s08.source_line.endswith("confidence: interview-anecdote")


def test_no_scorecard_makes_s08_placeholder():
    deck = build_deck(SPEC, RUN3)
    s08 = next(s for s in deck.slides if s.id == "S08")
    assert s08.kind == "placeholder" and s08.placeholder_reason == "requires engagement data"


# ------------------------------------------------------------------ leakage
def test_fingerprints_ignore_years_and_bare_small_ints():
    fp = withheld_fingerprints("AIA Group (whole) total assets US$345,423 million at 31 December 2025")
    assert fp["numbers"] == {"345423"} and fp["context_numbers"] == set()
    fp = withheld_fingerprints("Singlife total assets over S$16 billion")
    assert fp["context_numbers"] == {"16"}


def test_find_leaks_text_and_numbers():
    w = {"c1": "Prudential Assurance Co Singapore (I527L) S$66.3 billion under management as at 31 December 2025"}
    assert find_leaks("Peer size: Prudential Assurance Co Singapore (I527L) — withheld", w) == []
    assert [l["kind"] for l in find_leaks("AUM was 66.3 bn", w)] == ["number"]
    assert find_leaks("Prudential Assurance Co Singapore (I527L) S$66.3 billion under management", w)[0]["kind"] == "text"
    assert find_leaks("over S$16 billion", {"c": "assets over S$16 billion"})
    assert not find_leaks("page 16 of the report", {"c": "assets over S$16 billion"})


def test_lint_detects_withheld_leak_in_markdown_and_pptx(tmp):
    out = tmp / "o"
    write_outputs(build_deck(SPEC, RUN3), out)
    md = out / "deck-DRAFT.md"
    md.write_text(md.read_text().replace("## Four claims were removed", "Prudential manages S$66.3 billion.\n\n## Four claims were removed"))
    rep = lint_dir(out)
    leaks = [e for e in rep["errors"] if e["check"] == "withheld_leak"]
    assert leaks and "G-1" in leaks[0]["message"] and leaks[0]["message"].startswith("markdown:")
    # pptx surface too
    write_outputs(build_deck(SPEC, RUN3), out)
    prs = Presentation(str(out / "deck-DRAFT.pptx"))
    prs.slides[3].shapes.add_textbox(0, 0, 100, 100).text_frame.text = "AIA Group total assets US$345,423 million"
    prs.save(str(out / "deck-DRAFT.pptx"))
    rep = lint_dir(out)
    assert any(e["check"] == "withheld_leak" and e["message"].startswith("pptx:") for e in rep["errors"])


def test_client_profile_withholds_annual_report_and_nothing_leaks(tmp):
    """Client build, attachment pending: all 14 claims withheld (spec C2a)."""
    def withhold_all(rep):
        v = rep["verification"]
        for d in v["decisions"]:
            d["blocked"] = True
            d["status"] = "blocked"
            d["withheld_awaiting_reconciliation"] = d["claim_id"].startswith("great-eastern.ar.")
        rep["released_claim_ids"] = []
        for s in rep["signals"]:
            if s.get("evidence_claim_ids"):
                s["status"] = "blocked"
    run = make_run(tmp, recon=None, resolution=0.2, mutate=withhold_all)
    deck = build_deck(SPEC, run)
    rep = write_outputs(deck, tmp / "o")
    assert rep["ok"], rep["errors"]
    s09 = next(s for s in deck.slides if s.id == "S09")
    assert s09.title == "Great Eastern's outside-in AI signals are partly withheld pending verification"
    a1 = next(s for s in deck.slides if s.id == "A1")
    assert a1.kind == "placeholder" and a1.placeholder_reason == "withheld"
    md = (tmp / "o" / "deck-DRAFT.md").read_text()
    assert "75%" not in md and "Zhao" not in md.replace("ZHAO", "")
    assert "awaiting publisher reconciliation" in md
    assert "Released: 0 claims. Withheld: 14 claims (listed)." in md


# ------------------------------------------------------------------- lint rules
def _spec_with(tmp, old, new):
    p = tmp / "spec.yaml"
    p.write_text(SPEC.read_text().replace(old, new))
    return p


def test_title_over_15_words_is_lint_error(tmp):
    long = '"{client}\'s FY{fy} annual report reaches level {rubric_level} of 5 on the AI-disclosure rubric, which corroborates disclosure only"'
    spec = _spec_with(tmp, '"{client}\'s FY{fy} annual report reaches level {rubric_level} of 5 on the AI-disclosure rubric"', long)
    rep = write_outputs(build_deck(spec, RUN3), tmp / "o")
    errs = [e for e in rep["errors"] if e["check"] == "title_length"]
    assert len(errs) == 1 and errs[0]["slide"] == "S09" and "18 words" in errs[0]["message"]


def test_missing_source_line_is_lint_error(tmp):
    out = tmp / "o"
    write_outputs(build_deck(SPEC, RUN3), out)
    md = out / "deck-DRAFT.md"
    md.write_text(re.sub(r"^Source: T2 .*$", "", md.read_text(), flags=re.M))
    rep = lint_dir(out)
    assert {e["slide"] for e in rep["errors"] if e["check"] == "source_line"} == {"S09", "A1"}


def test_banner_removed_is_lint_error(tmp):
    out = tmp / "o"
    write_outputs(build_deck(SPEC, RUN3), out)
    md = out / "deck-DRAFT.md"
    md.write_text(md.read_text().replace(f"> **{BANNER} ({GEH_REASONS})**", "", 1))
    rep = lint_dir(out)
    assert [e["slide"] for e in rep["errors"] if e["check"] == "banner"] == ["S01"]


def test_lint_rechecks_build_report_truth(tmp):
    """A sign-off added after the build does not silently upgrade the deck; and a
    deck built final then re-linted against a now-failing report is flagged."""
    run = make_run(tmp)
    so = make_signoff(tmp, run)
    out = tmp / "o"
    write_outputs(build_deck(SPEC, run, signoff=so), out)
    rep_path = run / "report.json"
    rep = json.loads(rep_path.read_text())
    rep["verification"]["metrics"]["citation_resolution_rate"] = 0.9
    rep_path.write_text(json.dumps(rep))
    lint = lint_dir(out)
    checks = {e["check"] for e in lint["errors"]}
    assert {"interlock_drift", "draft_suffix", "banner", "s02_statement"} <= checks


def test_s02_truth_checked(tmp):
    out = tmp / "o"
    write_outputs(build_deck(SPEC, RUN3), out)
    md = out / "deck-DRAFT.md"
    md.write_text(md.read_text().replace("Released: 10 claims. Withheld: 4 claims (listed).", "Released: 14 claims. Withheld: 0 claims (listed)."))
    rep = lint_dir(out)
    assert any(e["check"] == "s02_truth" for e in rep["errors"])


# ------------------------------------------------------------ no override path
def test_cli_has_no_force_flag(tmp, capsys):
    with pytest.raises(SystemExit) as e:
        main(["build", "--spec", str(SPEC), "--run", str(RUN3), "--out", str(tmp / "o"), "--force"])
    assert e.value.code == 2
    with pytest.raises(SystemExit):
        main(["build", "--spec", str(SPEC), "--run", str(RUN3), "--out", str(tmp / "o"), "--fo"])  # no abbreviations


def test_spec_cannot_carry_override_keys(tmp):
    for bad in ("force: true\n", "interlock: off\n"):
        p = tmp / "s.yaml"
        p.write_text(bad + SPEC.read_text())
        with pytest.raises(SpecError):
            load_spec(p)
    p = tmp / "s2.yaml"
    p.write_text(SPEC.read_text().replace("  - id: S09\n", "  - id: S09\n    force_final: true\n"))
    with pytest.raises(SpecError):
        load_spec(p)


def test_cli_build_and_lint(tmp, capsys):
    rc = main(["build", "--spec", str(SPEC), "--run", str(RUN3), "--out", str(tmp / "o")])
    out = capsys.readouterr().out
    assert rc == 0 and "Released: 10 claims. Withheld: 4 claims (listed)." in out and GEH_REASONS in out
    assert main(["lint", str(tmp / "o")]) == 0
    assert json.loads((tmp / "o" / "lint.json").read_text())["ok"]


# ------------------------------------------------------------------- run 4
@pytest.mark.skipif(not (RUN4 / "report.json").exists(), reason="run 4 not present")
def test_run4_parent_context_signal_not_scored(tmp):
    deck = build_deck(SPEC, RUN4)
    rep = write_outputs(deck, tmp / "o")
    assert rep["ok"], rep["errors"]
    md = (tmp / "o" / "deck-DRAFT.md").read_text()
    row = next(l for l in md.splitlines() if l.startswith("| Parent-group"))
    assert "D1.2 context — not scored" in row and "context only" in row
