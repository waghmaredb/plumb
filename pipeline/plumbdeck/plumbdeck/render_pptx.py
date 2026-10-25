"""PPTX renderer (python-pptx), 16:9, neutral template: no client branding, no logos.

Named shapes the lint pass reads back:
  plumbdeck-banner   DRAFT banner (present on every slide when the interlock fails)
  plumbdeck-title    action title
  plumbdeck-source   source line
  plumbdeck-footer   §12.2 short form / draft footer
Each slide's notes start with ``plumbdeck:<ID>`` (``plumbdeck:A1 2/2`` on continuations).
"""
from __future__ import annotations

import math
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .model import Bars, Block, Boxes, Bullets, Cells, Columns, Para, Slide, Table

W, H = 13.333, 7.5
ML, MR = 0.6, 0.6
FONT = "Arial"
INK = RGBColor(0x1F, 0x23, 0x28)
MUTED = RGBColor(0x57, 0x60, 0x6A)
RULE = RGBColor(0xD0, 0xD7, 0xDE)
ACCENT = RGBColor(0x2F, 0x4B, 0x66)
ACCENT_LIGHT = RGBColor(0xA9, 0xBC, 0xCF)
TRACK = RGBColor(0xEC, 0xEF, 0xF2)
HEAD_FILL = RGBColor(0xEE, 0xF1, 0xF4)
BANNER_FILL = RGBColor(0xB4, 0x23, 0x18)
WITHHELD_INK = RGBColor(0x8A, 0x4B, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _lines(text: str, width_in: float, pt: float) -> int:
    chars_per_line = max(8, int(width_in / (pt * 0.52 / 72)))
    n = 0
    for part in str(text).split("\n"):
        n += max(1, math.ceil(len(part) / chars_per_line))
    return n


def _text(slide, x, y, w, h, text, pt=12, bold=False, color=INK, italic=False, name=None, align=None,
          anchor=MSO_ANCHOR.TOP, fill=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    if fill is not None:
        tb.fill.solid()
        tb.fill.fore_color.rgb = fill
    p = tf.paragraphs[0]
    if align:
        p.alignment = align
    r = p.add_run()
    r.text = str(text)
    f = r.font
    f.name, f.size, f.bold, f.italic = FONT, Pt(pt), bold, italic
    f.color.rgb = color
    return tb


def _rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(max(w, 0.001)), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    st = s._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if st is not None:  # drop theme style refs (LibreOffice draws effectRef shadows otherwise)
        s._element.remove(st)
    return s


def _line(slide, x1, y1, x2, y2, color=RULE, width_pt=0.75):
    """Axis-aligned rule drawn as a thin borderless rectangle (connectors pick up theme shadows)."""
    t = width_pt / 72
    if abs(y2 - y1) < 1e-6:
        return _rect(slide, min(x1, x2), y1 - t / 2, abs(x2 - x1), t, color)
    return _rect(slide, x1 - t / 2, min(y1, y2), t, abs(y2 - y1), color)


# ------------------------------------------------------------------ blocks
def _para_h(b: Para, w: float) -> float:
    pt = {"lead": 14, "strong": 12, "muted": 10, "withheld": 14, "withheld_inline": 11}.get(b.style, 11)
    return _lines(b.text, w, pt * 0.95) * pt * 1.25 / 72 + 0.08


def _row_heights(t: Table, widths_in: list[float]) -> list[float]:
    lh = t.font_pt * 1.22 / 72
    hs = []
    for r in [t.header] + t.rows:
        n = max(_lines(c, w - 0.12, t.font_pt * 0.97) for c, w in zip(r, widths_in))
        hs.append(n * lh + 0.12)  # LibreOffice / PowerPoint add paragraph spacing per cell; estimate high, never overlap
    return hs


def _col_widths(t: Table, w: float) -> list[float]:
    rel = t.widths or [1.0] * len(t.header)
    s = sum(rel)
    return [w * r / s for r in rel]


BOX_GAP, BOX_PAD, BOX_LABEL_H = 0.15, 0.1, 0.3


def _boxes_row_heights(b: Boxes, w: float) -> list[float]:
    cols = max(1, b.columns)
    bw = (w - BOX_GAP * (cols - 1)) / cols
    hs = []
    for r in range(0, len(b.boxes), cols):
        row = b.boxes[r:r + cols]
        hs.append(max(BOX_LABEL_H + sum(_para_h(p, bw - 2 * BOX_PAD) for p in paras) + 2 * BOX_PAD for _, paras in row) + BOX_GAP)
    return hs


def _block_h(b: Block, w: float) -> float:
    if isinstance(b, Para):
        return _para_h(b, w)
    if isinstance(b, Bullets):
        return sum(_lines("• " + i, w, 11) * 11 * 1.3 / 72 + 0.04 for i in b.items) + 0.1
    if isinstance(b, Table):
        return sum(_row_heights(b, _col_widths(b, w))) + 0.15
    if isinstance(b, Bars):
        return 0.55 * len(b.bars) + 0.45
    if isinstance(b, Cells):
        return math.ceil(len(b.cells) / b.columns) * 0.95 + 0.1
    if isinstance(b, Boxes):
        return sum(_boxes_row_heights(b, w)) + 0.1
    if isinstance(b, Columns):
        return max(sum(_block_h(x, w * b.ratio - 0.15) for x in b.left),
                   sum(_block_h(x, w * (1 - b.ratio) - 0.15) for x in b.right))
    return 0.3


def _draw_block(slide, b: Block, x: float, y: float, w: float) -> float:
    if isinstance(b, Para):
        style = {"lead": (14, True, INK, False), "strong": (12, True, INK, False), "muted": (10, False, MUTED, False),
                 "withheld": (14, True, WITHHELD_INK, False),
                 "withheld_inline": (11, True, WITHHELD_INK, False)}.get(b.style, (11, False, INK, False))
        h = _para_h(b, w)
        _text(slide, x, y, w, h, b.text, pt=style[0], bold=style[1], color=style[2], italic=style[3])
        return h
    if isinstance(b, Bullets):
        h = _block_h(b, w)
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, item in enumerate(b.items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = f"• {item}"
            r.font.name, r.font.size = FONT, Pt(11)
            r.font.color.rgb = INK
            p.space_after = Pt(3)
        return h
    if isinstance(b, Table):
        cw = _col_widths(b, w)
        rh = _row_heights(b, cw)
        shape = slide.shapes.add_table(len(b.rows) + 1, len(b.header), Inches(x), Inches(y), Inches(w), Inches(sum(rh)))
        tbl = shape.table
        tbl.first_row = True
        for j, cwi in enumerate(cw):
            tbl.columns[j].width = Emu(int(Inches(cwi)))
        for i, r in enumerate([b.header] + b.rows):
            tbl.rows[i].height = Emu(int(Inches(rh[i])))
            for j, val in enumerate(r):
                cell = tbl.cell(i, j)
                cell.margin_left = cell.margin_right = Inches(0.05)
                cell.margin_top = cell.margin_bottom = Inches(0.02)
                cell.fill.solid()
                cell.fill.fore_color.rgb = HEAD_FILL if i == 0 else WHITE
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                run = p.add_run()
                run.text = str(val)
                f = run.font
                f.name, f.size, f.bold = FONT, Pt(b.font_pt), i == 0
                f.color.rgb = WITHHELD_INK if str(val).startswith("Withheld") else INK
        return sum(rh) + 0.15
    if isinstance(b, Bars):
        label_w, val_w = 2.3, 1.6
        track_w = w - label_w - val_w
        yy = y + 0.05
        colors = [ACCENT_LIGHT, ACCENT]
        for k, bar in enumerate(b.bars):
            _text(slide, x, yy, label_w, 0.4, bar.label, pt=11, bold=True, anchor=MSO_ANCHOR.MIDDLE)
            _rect(slide, x + label_w, yy + 0.08, track_w, 0.26, TRACK)
            _rect(slide, x + label_w, yy + 0.08, track_w * max(0.0, min(100.0, bar.value)) / 100, 0.26,
                  colors[min(k, 1)] if len(b.bars) > 1 else ACCENT)
            _text(slide, x + label_w + track_w + 0.05, yy, val_w, 0.4, f"{bar.value:.1f}  {bar.note}", pt=11,
                  anchor=MSO_ANCHOR.MIDDLE)
            yy += 0.55
        for v, lbl in b.ticks:
            tx = x + label_w + track_w * v / 100
            _line(slide, tx, y + 0.05, tx, yy - 0.1, RULE, 0.5)
            _text(slide, tx - 1.0, yy - 0.08, 1.0, 0.25, lbl, pt=8, color=MUTED, align=PP_ALIGN.RIGHT)
        if b.boundary is not None:
            bx = x + label_w + track_w * b.boundary / 100
            ln = _line(slide, bx, y, bx, yy - 0.05, BANNER_FILL, 1.5)
            _text(slide, bx + 0.03, y - 0.2, 2.0, 0.22, f"{b.boundary_label} ({b.boundary:g})", pt=8, color=BANNER_FILL)
        return 0.55 * len(b.bars) + 0.45
    if isinstance(b, Cells):
        cols = b.columns
        cw = (w - 0.1 * (cols - 1)) / cols
        for i, (lbl, val, note) in enumerate(b.cells):
            cx = x + (i % cols) * (cw + 0.1)
            cy = y + (i // cols) * 0.95
            _rect(slide, cx, cy, cw, 0.85, HEAD_FILL)
            _text(slide, cx + 0.05, cy + 0.03, cw - 0.1, 0.25, lbl, pt=9, bold=True, color=MUTED)
            _text(slide, cx + 0.05, cy + 0.22, cw - 0.1, 0.32, val, pt=14, bold=True, color=ACCENT)
            _text(slide, cx + 0.05, cy + 0.52, cw - 0.1, 0.33, note[:90], pt=7, color=MUTED)
        return math.ceil(len(b.cells) / cols) * 0.95 + 0.1
    if isinstance(b, Boxes):
        cols = max(1, b.columns)
        bw = (w - BOX_GAP * (cols - 1)) / cols
        yy = y
        for r, rh in zip(range(0, len(b.boxes), cols), _boxes_row_heights(b, w)):
            for k, (label, paras) in enumerate(b.boxes[r:r + cols]):
                bx = x + k * (bw + BOX_GAP)
                box = _rect(slide, bx, yy, bw, rh - BOX_GAP, HEAD_FILL)
                _rect(slide, bx, yy, 0.06, rh - BOX_GAP, ACCENT)
                _text(slide, bx + BOX_PAD, yy + BOX_PAD - 0.04, bw - 2 * BOX_PAD, BOX_LABEL_H, label, pt=11, bold=True, color=ACCENT)
                py = yy + BOX_PAD + BOX_LABEL_H
                for p in paras:
                    py += _draw_block(slide, p, bx + BOX_PAD, py, bw - 2 * BOX_PAD)
            yy += rh
        return yy - y + 0.1
    if isinstance(b, Columns):
        lw = w * b.ratio - 0.15
        rw = w * (1 - b.ratio) - 0.15
        yl = y
        for blk in b.left:
            yl += _draw_block(slide, blk, x, yl, lw)
        yr = y
        for blk in b.right:
            yr += _draw_block(slide, blk, x + lw + 0.3, yr, rw)
        return max(yl, yr) - y
    raise TypeError(type(b))


# ------------------------------------------------------------------ pages
def _pages(s: Slide) -> list[list[Block]]:
    """Paginate an oversize top-level table across continuation slides."""
    for idx, b in enumerate(s.blocks):
        if isinstance(b, Table) and len(b.rows) > b.rows_per_slide:
            chunks = [b.rows[i:i + b.rows_per_slide] for i in range(0, len(b.rows), b.rows_per_slide)]
            pages = []
            for k, ch in enumerate(chunks):
                t = Table(b.header, ch, b.widths, b.font_pt, b.rows_per_slide)
                if k == 0:
                    pages.append(s.blocks[:idx] + [t] + s.blocks[idx + 1:])
                else:
                    pages.append([t])
            return pages
    return [s.blocks]


def _draw_slide(prs, s: Slide, blocks: list[Block], page: int, npages: int, number: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    top = 0.35
    if s.banner:
        _text(slide, 0, 0, W, 0.42, s.banner, pt=11, bold=True, color=WHITE, name="plumbdeck-banner",
              fill=BANNER_FILL, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
        top = 0.55
    cont = f"  ({page}/{npages})" if npages > 1 else ""
    tag = f"  ·  {s.placeholder_reason}" if s.placeholder_reason else ""
    _text(slide, ML, top, W - ML - MR, 0.3, f"{s.id} · {s.section}{cont}{tag}", pt=10, color=MUTED, name="plumbdeck-section")
    body_w = W - ML - MR
    if s.kind == "cover":
        _text(slide, ML, top + 1.4, body_w, 1.4, s.title, pt=30, bold=True, name="plumbdeck-title")
        _line(slide, ML, top + 2.95, ML + 3.0, top + 2.95, ACCENT, 2.0)
        y = top + 3.15
    else:
        tl = min(2, _lines(s.title, body_w - 0.3, 20))
        th = tl * 0.36 + 0.15
        _text(slide, ML, top + 0.35, body_w, th, s.title, pt=20, bold=True, name="plumbdeck-title",
              anchor=MSO_ANCHOR.TOP)
        _line(slide, ML, top + 0.4 + th, W - MR, top + 0.4 + th)
        y = top + 0.55 + th
    for b in blocks:
        y += _draw_block(slide, b, ML, y, body_w)
    if s.so_what and page == 1:
        sy = min(max(y + 0.05, 5.7), 6.0)
        box = _text(slide, ML, sy, body_w, 0.45, f"So what: {s.so_what}", pt=11, bold=True, color=INK,
                    fill=HEAD_FILL, anchor=MSO_ANCHOR.MIDDLE, name="plumbdeck-sowhat")
        box.line.color.rgb = ACCENT
    if s.source_line:
        _text(slide, ML, 6.5, body_w, 0.38, s.source_line, pt=9, color=MUTED, name="plumbdeck-source")
    _line(slide, ML, 6.92, W - MR, 6.92)
    _text(slide, ML, 6.97, body_w - 0.8, 0.4, s.footer, pt=8, color=MUTED, name="plumbdeck-footer")
    _text(slide, W - MR - 0.7, 6.97, 0.7, 0.3, str(number), pt=8, color=MUTED, align=PP_ALIGN.RIGHT)
    slide.notes_slide.notes_text_frame.text = f"plumbdeck:{s.id}" + (f" {page}/{npages}" if npages > 1 else "") + (
        f"\n{s.placeholder_reason}" if s.placeholder_reason else "")


def render_pptx(deck, path: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = f"{deck.run.client_full} — {deck.spec.name}" + (" (DRAFT)" if deck.draft else "")
    prs.core_properties.author = "plumbdeck"
    prs.core_properties.subject = deck.interlock.banner or deck.interlock.outcome_sentence
    n = 0
    for s in deck.slides:
        pages = _pages(s)
        for k, blocks in enumerate(pages, 1):
            n += 1
            _draw_slide(prs, s, blocks, k, len(pages), n)
    prs.save(str(path))
    return Path(path)
