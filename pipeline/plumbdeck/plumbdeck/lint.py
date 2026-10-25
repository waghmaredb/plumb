"""Deck-export lint (pipeline step S8; ghost deck R3, R4, R6, R8, R9).

``lint_dir`` re-reads the build report from the run folder and recomputes the
interlock (truth, not presence), then checks the rendered Markdown, the PPTX
(read back with python-pptx) and titles.md. Errors fail the lint; warnings do not.
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Optional

from .interlock import BANNER, compute_interlock
from .runio import RunData, RunError

MAX_TITLE_WORDS = 15
SOURCE_RE = re.compile(
    r"^Source: T[1-5] .+?, \d{4}-\d{2}-\d{2}(?:; T[1-5] .+?, \d{4}-\d{2}-\d{2})*; confidence: [a-z][a-z\- /]*$"
)
BUILD_SOURCE_RE = re.compile(r"^Source: PLUMB build report \S+, \d{4}-\d{2}-\d{2}$")
SLIDE_MARK = re.compile(r"^<!-- plumbdeck:slide id=(\S+) kind=(\S+) -->$", re.M)
NUM_RE = re.compile(
    r"(?<![A-Za-z0-9.,])(?P<cur>S\$|US\$|A\$|\$|€|£)?(?P<num>\d{1,3}(?:,\d{3})+|\d+)(?P<dec>\.\d+)?"
    r"(?P<post>\s?(?:%|per cent|billion|million|bn\b|m\b))?"
)
VERIFICATION_PARA = "Verification performed."


def title_words(t: str) -> int:
    return len([w for w in str(t).split() if re.search(r"[A-Za-z0-9]", w)])


# ------------------------------------------------------------ leak detection
def _words(s: str) -> list[str]:
    s = s.lower().replace("’", "'").replace("“", '"').replace("”", '"')
    return re.findall(r"[a-z0-9]+", s)


def _numbers(s: str) -> list[tuple[str, bool]]:
    """(normalised number, has_context) for every numeric token."""
    out = []
    for m in NUM_RE.finditer(s):
        val = m.group("num").replace(",", "") + (m.group("dec") or "")
        out.append((val, bool(m.group("cur") or m.group("post"))))
    return out


def withheld_fingerprints(text: str) -> dict[str, Any]:
    """Shingles of the claim's words plus its distinctive numbers.

    Distinctive = has a decimal part or three or more integer digits (years 1900-2100
    excepted); one- and two-digit integers count only with currency / % / scale
    context, and then only match a rendered number that also carries such context."""
    ws = _words(text)
    n = 8 if len(ws) >= 8 else len(ws)
    shingles = set()
    if n >= 4:
        shingles = {" ".join(ws[i:i + n]) for i in range(0, len(ws) - n + 1)}
    free, ctx = set(), set()
    for m in NUM_RE.finditer(text):
        ip, dec = m.group("num").replace(",", ""), m.group("dec") or ""
        val = ip + dec
        if not dec and len(ip) == 4 and 1900 <= int(ip) <= 2100:
            continue
        if dec or len(ip) >= 3:
            free.add(val)
        elif m.group("cur") or m.group("post"):
            ctx.add(val)
    return {"shingles": shingles, "numbers": free, "context_numbers": ctx}


def find_leaks(rendered: str, withheld: dict[str, str]) -> list[dict[str, str]]:
    joined = " " + " ".join(_words(rendered)) + " "
    nums = _numbers(rendered)
    free_nums = {v for v, _ in nums}
    ctx_nums = {v for v, c in nums if c}
    leaks = []
    for cid, text in withheld.items():
        fp = withheld_fingerprints(text)
        for sh in fp["shingles"]:
            if f" {sh} " in joined:
                leaks.append({"claim_id": cid, "kind": "text", "match": sh})
                break
        for v in sorted(fp["numbers"] & free_nums):
            leaks.append({"claim_id": cid, "kind": "number", "match": v})
        for v in sorted(fp["context_numbers"] & ctx_nums):
            leaks.append({"claim_id": cid, "kind": "number", "match": v})
    return leaks


# ------------------------------------------------------------------ parsing
def parse_markdown(md: str) -> list[dict[str, Any]]:
    marks = list(SLIDE_MARK.finditer(md))
    slides = []
    for i, m in enumerate(marks):
        body = md[m.end(): marks[i + 1].start() if i + 1 < len(marks) else len(md)]
        title = next((l[3:].strip() for l in body.splitlines() if l.startswith("## ")), "")
        sources = [l.strip() for l in body.splitlines() if l.startswith("Source: ")]
        banner = next((l for l in body.splitlines() if l.startswith(f"> **{BANNER}")), None)
        slides.append({"id": m.group(1), "kind": m.group(2), "title": title, "sources": sources,
                       "banner": banner, "text": body})
    return slides


def _shape_text(shape) -> str:
    parts = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        parts.append(shape.text_frame.text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def parse_pptx(path: Path) -> list[dict[str, Any]]:
    from pptx import Presentation

    prs = Presentation(str(path))
    out = []
    for i, sl in enumerate(prs.slides, 1):
        notes = sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ""
        m = re.match(r"plumbdeck:(\S+)(?: (\d+)/(\d+))?", notes or "")
        named = {sh.name: sh for sh in sl.shapes}
        out.append({
            "index": i,
            "id": m.group(1) if m else None,
            "page": int(m.group(2)) if m and m.group(2) else 1,
            "title": named["plumbdeck-title"].text_frame.text if "plumbdeck-title" in named else None,
            "banner": named["plumbdeck-banner"].text_frame.text if "plumbdeck-banner" in named else None,
            "source": named["plumbdeck-source"].text_frame.text if "plumbdeck-source" in named else None,
            "footer": named["plumbdeck-footer"].text_frame.text if "plumbdeck-footer" in named else None,
            "text": "\n".join(_shape_text(sh) for sh in sl.shapes),
            "width": prs.slide_width, "height": prs.slide_height,
        })
    return out


def parse_titles(md: str) -> list[tuple[str, str]]:
    rows = []
    for l in md.splitlines():
        m = re.match(r"^\| (S\d+|A\d+) \| .*? \| (.*) \| (?:cover \(R3 n/a\)|\d+) \| [^|]+ \|$", l)
        if m:
            rows.append((m.group(1), m.group(2).replace("\\|", "|")))
    return rows


# --------------------------------------------------------------------- lint
def lint_dir(out: Path, write: bool = True) -> dict[str, Any]:
    out = Path(out)
    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []

    def err(check, msg, slide=None):
        errors.append({"check": check, "slide": slide, "message": msg})

    def warn(check, msg, slide=None):
        warnings.append({"check": check, "slide": slide, "message": msg})

    man_path = out / "build.json"
    if not man_path.is_file():
        err("manifest", "build.json missing; run `plumbdeck build` first")
        return _finish(out, errors, warnings, None, write)
    man = json.loads(man_path.read_text(encoding="utf-8"))
    meta = {s["id"]: s for s in man["slides"]}

    # ---- truth: recompute the interlock from the build report
    try:
        run = RunData(Path(man["run_dir"]))
        il = compute_interlock(run, Path(man["signoff"]) if man.get("signoff") else None)
    except (RunError, OSError, ValueError) as e:
        err("build_report", f"cannot re-read the build report: {e}")
        return _finish(out, errors, warnings, None, write)
    if il.passed != man["interlock"]["passed"] or il.reasons != man["interlock"]["reasons"]:
        err("interlock_drift", f"deck was built with interlock {man['interlock']['reasons'] or 'passed'} but the build report now gives {il.reasons or 'passed'}; rebuild")
    draft = not il.passed

    md_p, pptx_p, titles_p = out / man["deck_md"], out / man["deck_pptx"], out / man["titles_md"]
    missing = [p for p in (md_p, pptx_p, titles_p) if not p.is_file()]
    for p in missing:
        err("files", f"{p.name} missing")
    if missing:
        return _finish(out, errors, warnings, il, write)

    # ---- file naming
    if draft:
        if not md_p.stem.endswith("-DRAFT") or not pptx_p.stem.endswith("-DRAFT"):
            err("draft_suffix", "interlock failed but the deck filename lacks the -DRAFT suffix")
        for stale in ("deck.md", "deck.pptx"):
            if (out / stale).exists():
                err("draft_suffix", f"{stale} (final name) present beside a draft build")
    else:
        if md_p.stem.endswith("-DRAFT") or pptx_p.stem.endswith("-DRAFT"):
            err("draft_suffix", "interlock passed but the deck carries the -DRAFT suffix; rebuild")

    md = md_p.read_text(encoding="utf-8")
    mds = parse_markdown(md)
    pps = parse_pptx(pptx_p)
    titles = parse_titles(titles_p.read_text(encoding="utf-8"))

    # ---- structure: md, pptx and titles agree
    md_ids = [s["id"] for s in mds]
    pp_ids = list(dict.fromkeys(s["id"] for s in pps))
    if md_ids != [s["id"] for s in man["slides"]]:
        err("structure", "deck Markdown slide order differs from the build manifest")
    if md_ids != pp_ids:
        err("structure", f"Markdown and PPTX slide IDs differ: {md_ids} vs {pp_ids}")
    if [t[0] for t in titles] != md_ids or [t[1] for t in titles] != [s["title"] for s in mds]:
        err("horizontal_flow", "titles.md does not match the deck's titles in order")
    pp_first = {s["id"]: s for s in pps if s["page"] == 1}
    for s in mds:
        p = pp_first.get(s["id"])
        if p and (p["title"] or "").strip() != s["title"]:
            err("structure", "PPTX title differs from Markdown title", s["id"])
    if "S02" not in md_ids:
        err("s02_missing", "no S02 disclosure slide (R9)")
    if abs(pps[0]["width"] / pps[0]["height"] - 16 / 9) > 0.01 if pps else True:
        err("aspect", "PPTX is not 16:9")

    # ---- titles (R3)
    for s in mds:
        m = meta.get(s["id"], {})
        if m.get("title_exempt"):
            continue
        n = title_words(s["title"])
        if n > MAX_TITLE_WORDS:
            err("title_length", f"action title has {n} words (max {MAX_TITLE_WORDS}): {s['title']}", s["id"])
        if s["kind"] != "placeholder" and ("[H]" in s["title"] or re.search(r"\[[^\]]+\]", s["title"])):
            (err if not draft else warn)("title_residue", f"title carries [H] / [bracket] residue: {s['title']}", s["id"])

    # ---- source lines (R6)
    for s in mds:
        m = meta.get(s["id"], {})
        if not m.get("is_data") or m.get("kind") == "placeholder":
            continue
        rx = BUILD_SOURCE_RE if m.get("kind") == "register" else SOURCE_RE
        if not s["sources"]:
            err("source_line", "data slide has no source line", s["id"])
        elif not all(rx.match(x) for x in s["sources"]):
            err("source_line", f"source line not in the required format: {s['sources'][0]}", s["id"])
        p = pp_first.get(s["id"])
        if p is not None and not p["source"]:
            err("source_line", "PPTX data slide has no source line", s["id"])
        if m.get("kind") == "data" and re.match(r"S(0[5-9]|[12]\d|30)$", s["id"]) and "**So what:**" not in s["text"]:
            warn("so_what", "evidence slide has no so-what line (R7)", s["id"])

    # ---- banner / S02 interlock (R8, R9)
    expected = il.banner
    for s in mds:
        if draft and (not s["banner"] or expected not in s["banner"]):
            err("banner", "Markdown slide lacks the DRAFT banner with the current reasons", s["id"])
        if not draft and s["banner"]:
            err("banner", "DRAFT banner present although the interlock passes; rebuild", s["id"])
    for p in pps:
        if draft and (not p["banner"] or p["banner"] != expected):
            err("banner", f"PPTX slide {p['index']} lacks the DRAFT banner with the current reasons", p["id"])
        if not draft and p["banner"]:
            err("banner", f"PPTX slide {p['index']} carries a DRAFT banner although the interlock passes", p["id"])
        if draft and p["footer"] and "All citations fetched and verified" in p["footer"]:
            err("footer", "short-form verification footer rendered on a draft", p["id"])
    s02 = next((s for s in mds if s["id"] == "S02"), None)
    s02p = [p for p in pps if p["id"] == "S02"]
    if s02:
        has_stmt = VERIFICATION_PARA in s02["text"] or any(VERIFICATION_PARA in p["text"] for p in s02p)
        if draft and has_stmt:
            err("s02_statement", "verification statement rendered although the interlock fails", "S02")
        if not draft and not has_stmt:
            err("s02_statement", "interlock passes but S02 lacks the verification statement", "S02")
        if il.outcome_sentence not in s02["text"]:
            err("s02_truth", f"S02 does not state the build report's outcome '{il.outcome_sentence}'", "S02")
        if il.resolution is not None and f"{il.resolution:g}" not in s02["text"]:
            err("s02_truth", "S02 resolution figure differs from the build report", "S02")
        if not draft and "[approver]" in s02["text"]:
            err("s02_truth", "S02 approver field is empty", "S02")

    # ---- withheld-claim leakage (R8, C11)
    withheld = {cid: run.claims[cid].text for cid in run.withheld_ids}
    surfaces = {"markdown": md, "pptx": "\n".join(p["text"] for p in pps), "titles": titles_p.read_text(encoding="utf-8")}
    for name, text in surfaces.items():
        for lk in find_leaks(text, withheld):
            err("withheld_leak", f"{name}: withheld claim {lk['claim_id']} ({run.claims[lk['claim_id']].gap_id}) leaks its {lk['kind']} '{lk['match']}'")

    # ---- state summary
    n_ph = sum(1 for s in mds if s["kind"] == "placeholder")
    if n_ph:
        ids = ", ".join(s["id"] for s in mds if s["kind"] == "placeholder")
        (err if not draft else warn)("placeholders", f"{n_ph} placeholder slide(s) with [H] titles (requires engagement data / withheld): {ids}; deck is not board-ready")
    return _finish(out, errors, warnings, il, write)


def _finish(out: Path, errors, warnings, il, write: bool) -> dict[str, Any]:
    rep = {
        "ok": not errors,
        "errors": errors,
        "warnings": warnings,
        "draft": None if il is None else not il.passed,
        "interlock": None if il is None else il.to_dict(),
    }
    if write:
        (Path(out) / "lint.json").write_text(json.dumps(rep, indent=2, ensure_ascii=False), encoding="utf-8")
    return rep
